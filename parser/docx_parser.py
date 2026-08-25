#!/usr/bin/env python3
"""
Парсер расписания из DOCX-файлов с записью в новую схему БД.

Что делает модуль:
- читает первую таблицу из .docx;
- парсит дату, диапазон времени, предмет, преподавателя и аудиторию;
- возвращает список словарей в старом удобном формате;
- записывает данные в новую БД через async_session:
  Group, Subject, Teacher, Classroom, Lesson, LessonGroup.

Важно:
- поддерживается только .docx;
- функционал .doc -> .docx намеренно не включён.
"""

from __future__ import annotations

import argparse
import asyncio
import re
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Optional

from docx import Document
from sqlalchemy import select

from database.connection import async_session
from database.models import Classroom, Group, Lesson, LessonGroup, Subject, Teacher
from parser.regex_patterns import (
    ACADEMIC_TITLES_PATTERN,
    CLEAN_NAME_PATTERN,
    DATE_PATTERN,
)
from utils.timezone import YEKT_TZ


TIME_RANGE_PATTERN = re.compile(
    r"\b(\d{1,2})[:.](\d{2})\s*[-–—]\s*(\d{1,2})[:.](\d{2})\b"
)

TIME_START_PATTERN = re.compile(r"\b(\d{1,2})[:.](\d{2})\b")

EMPTY_CLASSROOM_VALUES = {
    "",
    "-",
    "—",
    "–",
    "не указана",
    "не указано",
    "ауд. не указана",
    "аудитория не указана",
}


def validate_schedule_file(file_path: str) -> bool:
    """
    Проверяет, что файл существует и это .docx.
    """
    path = Path(file_path)

    if not path.exists():
        print(f"❌ Файл не найден: {file_path}")
        return False

    if path.suffix.lower() != ".docx":
        print(f"❌ Поддерживается только .docx: {path.suffix}")
        return False

    return True


def clean_and_split_row(line: str) -> list[str]:
    """
    Разбивает строку таблицы по символу ячейки.
    """
    return [part.strip() for part in line.split("\x07")]


def parse_date(text: str) -> Optional[date]:
    """
    Ищет дату в формате ДД.ММ.ГГ или ДД.ММ.ГГГГ, также допускается дефис.
    """
    match = DATE_PATTERN.search(text)
    if not match:
        return None

    day, month, year = match.groups()
    if len(year) == 2:
        year = "20" + year

    try:
        return date(int(year), int(month), int(day))
    except ValueError:
        return None


def parse_time_range(text: str) -> Optional[tuple[time, time]]:
    """
    Парсит диапазон времени вида:
    09.00-10.20
    09:00-10:20
    09.00 — 10.20
    """
    match = TIME_RANGE_PATTERN.search(text.replace(" ", ""))
    if not match:
        return None

    h1, m1, h2, m2 = match.groups()
    try:
        return time(int(h1), int(m1)), time(int(h2), int(m2))
    except ValueError:
        return None


def parse_start_time(text: str) -> Optional[time]:
    """
    Запасной вариант: парсит только старт времени.
    """
    match = TIME_START_PATTERN.search(text)
    if not match:
        return None

    hours, minutes = match.groups()
    try:
        return time(int(hours), int(minutes))
    except ValueError:
        return None


def determine_lesson_type(subject_text: str) -> str:
    """
    Определяет тип занятия по маркерам.
    """
    text_lower = subject_text.lower()

    if "(л)" in text_lower or "лекция" in text_lower:
        return "лекция"
    if "(пр)" in text_lower or "практика" in text_lower:
        return "практика"
    if "(конс)" in text_lower or "консультация" in text_lower:
        return "консультация"
    if "зачёт с оценкой" in text_lower or "зачет с оценкой" in text_lower or "диф.зачет" in text_lower:
        return "зачет с оценкой"
    if "зачёт" in text_lower or "зачет" in text_lower:
        return "зачет"
    if "экзамен" in text_lower:
        return "экзамен"

    return "другое"


def clean_subject_name(text: str) -> str:
    """
    Очищает название предмета от служебных пометок.
    """
    text = re.sub(r"\s*\((л|пр|конс)\)\s*", " ", text, flags=re.IGNORECASE)
    text = re.sub(
        r"\b(зач[её]т( с оценкой)?|диф\.?\s*зач[её]т)\b",
        " ",
        text,
        flags=re.IGNORECASE,
    )
    text = re.sub(r"\s+", " ", text).strip()
    return text.strip(",. -()")


def clean_teacher_name(text: str) -> Optional[str]:
    """
    Убирает должности и оставляет ФИО.
    """
    if not text:
        return None

    normalized = ACADEMIC_TITLES_PATTERN.sub("", text)
    normalized = re.sub(r"\s+", " ", normalized).strip(" ,.;:-")

    if not normalized:
        return None

    match = CLEAN_NAME_PATTERN.search(normalized)
    if match:
        return match.group(0).strip()

    return normalized


def normalize_classroom_name(text: str) -> Optional[str]:
    """
    Приводит аудиторию к нормальному виду.
    """
    if not text:
        return None

    value = re.sub(r"\s+", " ", text).strip(" ,.;:-")
    if value.lower() in EMPTY_CLASSROOM_VALUES:
        return None

    return value


def make_timezone_aware(current_date: date, lesson_time: time) -> datetime:
    """
    Собирает timezone-aware datetime.
    """
    return datetime.combine(current_date, lesson_time).replace(tzinfo=YEKT_TZ)


def split_subject(cell: str) -> tuple[str, Optional[str]]:
    """
    Старый удобный формат: отдельно предмет и тип.
    """
    lesson_type = determine_lesson_type(cell)
    subject = clean_subject_name(cell)
    return subject, lesson_type


def extract_fio(teacher_str: str) -> Optional[str]:
    """
    Извлекает ФИО преподавателя.
    """
    return clean_teacher_name(teacher_str)


def parse_docx_schedule(docx_path: str, group_name: str) -> list[dict[str, Any]]:
    """
    Парсит расписание из .docx файла и возвращает список словарей.

    Формат словаря:
    {
        "group": str,
        "subject": str,
        "lesson_type": str,
        "teacher": str | None,
        "classroom": str | None,
        "start_datetime": datetime,
        "end_datetime": datetime | None,
    }
    """
    doc = Document(docx_path)

    if not doc.tables:
        raise ValueError("В документе нет таблиц")

    table = doc.tables[0]
    lessons: list[dict[str, Any]] = []

    current_date: Optional[date] = None

    # Сохраняем старую логику: работаем по строкам первой таблицы
    # и обычно пропускаем заголовок.
    for row_idx, row in enumerate(table.rows[1:], start=2):
        cells = [cell.text.strip() for cell in row.cells]

        while len(cells) < 5:
            cells.append("")

        date_cell, time_cell, subject_cell, teacher_cell, room_cell = cells[:5]

        if date_cell:
            parsed_date = parse_date(date_cell)
            if parsed_date:
                current_date = parsed_date

        if not current_date:
            continue

        if not time_cell or not subject_cell:
            continue

        parsed_range = parse_time_range(time_cell)
        if parsed_range:
            t_start, t_end = parsed_range
        else:
            t_start = parse_start_time(time_cell)
            t_end = None
            if not t_start:
                continue

        subject, lesson_type = split_subject(subject_cell)

        if not subject or subject in {"-", "—", "–"}:
            continue

        teacher_name = extract_fio(teacher_cell) if teacher_cell else None
        classroom_name = normalize_classroom_name(room_cell)

        try:
            start_dt = make_timezone_aware(current_date, t_start)
            end_dt = make_timezone_aware(current_date, t_end) if t_end else None
        except ValueError:
            print(f"⚠️ Пропущена строка {row_idx}: некорректное время")
            continue

        lessons.append(
            {
                "group": group_name,
                "subject": subject,
                "lesson_type": lesson_type or "другое",
                "teacher": teacher_name,
                "classroom": classroom_name,
                "start_datetime": start_dt,
                "end_datetime": end_dt,
            }
        )

    print(f"✅ Успешно распарсено занятий: {len(lessons)}")
    return lessons


async def get_or_create(session, model, field_name: str, value: str):
    """
    Универсально получает объект или создаёт новый.
    """
    field = getattr(model, field_name)
    stmt = select(model).where(field == value)
    result = await session.execute(stmt)
    obj = result.scalars().first()

    if obj:
        return obj

    obj = model(**{field_name: value})
    session.add(obj)
    await session.flush()
    return obj


async def import_schedule_to_db_async(
    lessons: list[dict[str, Any]],
    group_name: str,
) -> int:
    """
    Импортирует список занятий в новую БД.
    """
    if not lessons:
        return 0

    imported_count = 0

    async with async_session() as session:
        group_obj = await get_or_create(session, Group, "name", group_name)

        for idx, lesson_data in enumerate(lessons, start=1):
            try:
                subject_name = lesson_data["subject"][:150].strip()
                if not subject_name:
                    print(f"⚠️ Пропущено занятие {idx}: пустой предмет")
                    continue

                subject_obj = await get_or_create(session, Subject, "name", subject_name)

                teacher_obj = None
                teacher_name = (lesson_data.get("teacher") or "").strip()
                if teacher_name:
                    teacher_obj = await get_or_create(session, Teacher, "name", teacher_name[:100])

                classroom_obj = None
                classroom_name = (lesson_data.get("classroom") or "").strip()
                if classroom_name:
                    classroom_obj = await get_or_create(
                        session,
                        Classroom,
                        "name",
                        classroom_name[:50],
                    )

                start_dt: datetime = lesson_data["start_datetime"]
                lesson_type: str = (lesson_data.get("lesson_type") or "другое").strip()

                conditions = [
                    Lesson.start_datetime == start_dt,
                    Lesson.subject_id == subject_obj.id,
                ]

                if classroom_obj is None:
                    conditions.append(Lesson.classroom_id.is_(None))
                else:
                    conditions.append(Lesson.classroom_id == classroom_obj.id)

                lesson_stmt = select(Lesson).where(*conditions)
                lesson_res = await session.execute(lesson_stmt)
                lesson_obj = lesson_res.scalars().first()

                if lesson_obj:
                    # Если занятие уже есть, аккуратно дополняем преподавателя, если его не было
                    if teacher_obj and lesson_obj.teacher_id is None:
                        lesson_obj.teacher_id = teacher_obj.id
                    if lesson_obj.type != lesson_type and lesson_type:
                        lesson_obj.type = lesson_type
                else:
                    lesson_obj = Lesson(
                        start_datetime=start_dt,
                        type=lesson_type,
                        classroom_id=classroom_obj.id if classroom_obj else None,
                        teacher_id=teacher_obj.id if teacher_obj else None,
                        subject_id=subject_obj.id,
                        comment=None,
                    )
                    session.add(lesson_obj)
                    await session.flush()

                link_stmt = select(LessonGroup).where(
                    LessonGroup.lesson_id == lesson_obj.id,
                    LessonGroup.group_id == group_obj.id,
                )
                link_res = await session.execute(link_stmt)
                if not link_res.scalars().first():
                    session.add(
                        LessonGroup(
                            lesson_id=lesson_obj.id,
                            group_id=group_obj.id,
                        )
                    )

                imported_count += 1

            except Exception as e:
                print(f"⚠️ Ошибка импорта занятия {idx}: {e}")
                print(f"   Данные: {lesson_data}")
                continue

        await session.commit()

    print(f"✅ Импорт завершён. Добавлено занятий: {imported_count}")
    return imported_count


async def parse_docx_to_db(file_path: str, group_name: str) -> int:
    """
    Полный цикл: парсинг DOCX и запись в БД.
    """
    if not validate_schedule_file(file_path):
        return 0

    lessons = parse_docx_schedule(file_path, group_name)
    return await import_schedule_to_db_async(lessons, group_name)


def import_schedule_to_db(lessons: list[dict[str, Any]], group_name: str) -> int:
    """
    Синхронная обёртка над асинхронным импортом.
    """
    return asyncio.run(import_schedule_to_db_async(lessons, group_name))


def print_schedule_summary(lessons: list[dict[str, Any]]) -> None:
    """
    Короткая сводка по распарсенным занятиям.
    """
    if not lessons:
        print("📋 Расписание пустое")
        return

    subjects = {lesson["subject"] for lesson in lessons}
    teachers = {lesson["teacher"] for lesson in lessons if lesson.get("teacher")}
    date_range = [lesson["start_datetime"] for lesson in lessons if lesson.get("start_datetime")]

    print("\n" + "=" * 50)
    print("📋 СВОДКА РАСПИСАНИЯ")
    print("=" * 50)
    print(f"Группа: {lessons[0]['group']}")
    print(f"Всего занятий: {len(lessons)}")
    print(f"Предметов: {len(subjects)}")
    print(f"Преподавателей: {len(teachers)}")

    if date_range:
        print(
            f"Период: {min(date_range).strftime('%d.%m.%Y')} - {max(date_range).strftime('%d.%m.%Y')}"
        )

    print("\nПредметы:")
    for subject in sorted(subjects):
        count = sum(1 for lesson in lessons if lesson["subject"] == subject)
        print(f"  • {subject}: {count} занятий")

    print("=" * 50 + "\n")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", type=str, required=True, help="Путь к .docx файлу")
    parser.add_argument("--group", type=str, required=True, help="Название группы")
    args = parser.parse_args()

    async def _run():
        lessons = parse_docx_schedule(args.file, args.group)
        print_schedule_summary(lessons)
        await import_schedule_to_db_async(lessons, args.group)

    asyncio.run(_run())


if __name__ == "__main__":
    main()
