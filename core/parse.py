import re
import asyncio

from pathlib import Path
from datetime import datetime

from docx import Document
from sqlalchemy import select

from core.database import AsyncSessionLocal

from core.models import (
    Teacher,
    Subject,
    Classroom,
    Group,
    Lesson,
    LessonGroup,
    LessonType,
)

# =========================================================
# CONFIG
# =========================================================

SCHEDULE_DIR = "./schedules"

# =========================================================
# REGEX
# =========================================================

DATE_REGEX = re.compile(
    r"(\d{2}\.\d{2}\.\d{2})"
)

TIME_REGEX = re.compile(
    r"(\d{1,2}\.\d{2}-\d{1,2}\.\d{2})"
)

LESSON_TYPE_REGEX = re.compile(
    r"\((л|пр|лаб|конс)\)"
)

# =========================================================
# HELPERS
# =========================================================

async def get_or_create(db, model, name: str):
    name = name.strip()

    stmt = select(model).where(
        model.name == name
    )

    result = await db.execute(stmt)

    obj = result.scalar_one_or_none()

    if obj:
        return obj

    obj = model(name=name)

    db.add(obj)

    await db.commit()
    await db.refresh(obj)

    print(f"[CREATE] {model.__name__}: {name}")

    return obj


def parse_lesson_type(subject_text: str):
    match = LESSON_TYPE_REGEX.search(subject_text)

    if not match:
        return LessonType.LECTURE

    value = match.group(1)

    if value == "л":
        return LessonType.LECTURE

    if value == "пр":
        return LessonType.PRACTICE

    if value == "лаб":
        return LessonType.LAB

    if value == "конс":
        return LessonType.PRACTICE

    return LessonType.LECTURE


def clean_subject(subject: str):
    subject = LESSON_TYPE_REGEX.sub("", subject)

    subject = (
        subject
        .replace("Зачёт.", "")
        .replace("Зачет.", "")
        .replace("Зачёт с оценкой.", "")
        .replace("Зачет с оценкой.", "")
    )

    return subject.strip(" .")


def parse_times(time_range: str):
    start_str, end_str = time_range.split("-")

    sh, sm = map(int, start_str.split("."))
    eh, em = map(int, end_str.split("."))

    return sh, sm, eh, em


# =========================================================
# DOCX
# =========================================================

def read_docx_lines(path: Path):
    doc = Document(path)

    lines = []

    # paragraphs
    for p in doc.paragraphs:
        text = p.text.strip()

        if text:
            lines.append(text)

    # tables
    for table in doc.tables:

        for row in table.rows:

            row_data = []

            for cell in row.cells:

                text = cell.text.strip()

                if text:
                    row_data.append(text)

            if row_data:
                lines.append(" ".join(row_data))

    cleaned = []

    for line in lines:

        line = (
            line
            .replace("\n", " ")
            .replace("\t", " ")
            .replace("￾", "")
        )

        line = re.sub(r"\s+", " ", line)

        line = line.strip()

        if line:
            cleaned.append(line)

    return cleaned


# =========================================================
# PARSER
# =========================================================

def split_lesson_line(line: str):
    """
    Пример строки:

    04.05.26 г Понедельник
    10.00-11.20 Эконометрика (л) к.п.н. Торсунова Э.Р. 410
    """

    # =============================================
    # DATE
    # =============================================

    date_match = DATE_REGEX.search(line)

    if not date_match:
        return None

    raw_date = date_match.group(1)

    try:
        current_date = datetime.strptime(
            raw_date,
            "%d.%m.%y"
        ).date()

    except:
        return None

    # =============================================
    # TIME
    # =============================================

    time_match = TIME_REGEX.search(line)

    if not time_match:
        return None

    time_range = time_match.group(1)

    # =============================================
    # REST
    # =============================================

    rest = line[time_match.end():].strip()

    parts = rest.split()

    if len(parts) < 3:
        return None

    # =============================================
    # CLASSROOM
    # =============================================

    if parts[-2] == "Ст.":
        classroom = "Ст. " + parts[-1]
        parts = parts[:-2]
    else:
        classroom = parts[-1]
        parts = parts[:-1]

    # =============================================
    # TEACHER
    # =============================================

    teacher_start = None

    for i in range(len(parts) - 1, -1, -1):

        if re.search(r"[А-ЯЁ]\.[А-ЯЁ]\.", parts[i]):
            teacher_start = i
            break

    if teacher_start is None:
        return None

    # захватываем должности
    while teacher_start > 0:

        prev = parts[teacher_start - 1]

        if (
            "." in prev
            or "," in prev
            or prev.lower() in [
                "доц.",
                "проф.",
                "асс.",
                "ст.",
            ]
        ):
            teacher_start -= 1
        else:
            break

    teacher = " ".join(parts[teacher_start:])

    subject = " ".join(parts[:teacher_start])

    return {
        "date": current_date,
        "time_range": time_range,
        "subject": subject.strip(),
        "teacher": teacher.strip(),
        "classroom": classroom.strip(),
    }


# =========================================================
# IMPORT
# =========================================================

async def import_schedule(path: Path):
    print(f"\n=== {path.name} ===")

    group_name = input(
        "Введите название группы: "
    ).strip()

    async with AsyncSessionLocal() as db:

        group = await get_or_create(
            db,
            Group,
            group_name
        )

        lines = read_docx_lines(path)

        imported = 0

        for line in lines:

            print(line)

            parsed = split_lesson_line(line)

            if not parsed:
                continue

            try:

                subject_name = clean_subject(
                    parsed["subject"]
                )

                teacher_name = parsed["teacher"]

                classroom_name = parsed["classroom"]

                lesson_type = parse_lesson_type(
                    parsed["subject"]
                )

                teacher = await get_or_create(
                    db,
                    Teacher,
                    teacher_name
                )

                subject = await get_or_create(
                    db,
                    Subject,
                    subject_name
                )

                classroom = await get_or_create(
                    db,
                    Classroom,
                    classroom_name
                )

                sh, sm, eh, em = parse_times(
                    parsed["time_range"]
                )

                current_date = parsed["date"]

                start_dt = datetime(
                    current_date.year,
                    current_date.month,
                    current_date.day,
                    sh,
                    sm,
                )

                end_dt = datetime(
                    current_date.year,
                    current_date.month,
                    current_date.day,
                    eh,
                    em,
                )

                # =========================================
                # DUPLICATES
                # =========================================

                stmt = (
                    select(Lesson)
                    .where(
                        Lesson.start_datetime == start_dt,
                        Lesson.end_datetime == end_dt,
                        Lesson.subject_id == subject.id,
                        Lesson.teacher_id == teacher.id,
                    )
                )

                result = await db.execute(stmt)

                lesson = result.scalar_one_or_none()

                if not lesson:

                    lesson = Lesson(
                        start_datetime=start_dt,
                        end_datetime=end_dt,
                        type=lesson_type,
                        classroom_id=classroom.id,
                        teacher_id=teacher.id,
                        subject_id=subject.id,
                        comment=""
                    )

                    db.add(lesson)

                    await db.commit()
                    await db.refresh(lesson)

                    print(
                        f"[LESSON CREATED] "
                        f"{subject_name}"
                    )

                # =========================================
                # GROUP LINK
                # =========================================

                link_stmt = (
                    select(LessonGroup)
                    .where(
                        LessonGroup.lesson_id == lesson.id,
                        LessonGroup.group_id == group.id,
                    )
                )

                link_result = await db.execute(
                    link_stmt
                )

                link = link_result.scalar_one_or_none()

                if not link:

                    db.add(
                        LessonGroup(
                            lesson_id=lesson.id,
                            group_id=group.id,
                        )
                    )

                    await db.commit()

                imported += 1

                print(
                    f"[OK] "
                    f"{parsed['date']} | "
                    f"{parsed['time_range']} | "
                    f"{subject_name}"
                )

            except Exception as e:

                print(
                    f"\n[ERROR]\n"
                    f"LINE: {line}\n"
                    f"ERROR: {e}\n"
                )

        print(
            f"\nИмпортировано занятий: {imported}"
        )


# =========================================================
# MAIN
# =========================================================

async def main():
    directory = Path(SCHEDULE_DIR)

    files = [
        *directory.glob("*.docx"),
        *directory.glob("*.doc"),
    ]

    if not files:
        print("Файлы не найдены")
        return

    print(f"Найдено файлов: {len(files)}")

    for file in files:

        try:
            await import_schedule(file)

        except Exception as e:

            print(
                f"\n[FATAL ERROR]\n"
                f"{file.name}\n"
                f"{e}"
            )


if __name__ == "__main__":
    asyncio.run(main())